"""Emballage des résultats validés : aucune modification du pipeline source."""
from __future__ import annotations

import json
import os
import tempfile
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Font
from PIL import Image

from . import contracts as c
from .adapters import Job
from .io import ContractError, digest, fmt, number, write_csv


def _jpeg(annotation: dict, out: Path, dpi: int) -> None:
    source = annotation["path"]
    if source.suffix.lower() == ".pdf":
        with pymupdf.open(source) as doc:
            index = annotation["page"] - 1
            if index >= len(doc):
                raise ContractError(f"{source.name}: page {index + 1} absente")
            page = doc[index]
            pix = page.get_pixmap(dpi=dpi, colorspace=pymupdf.csRGB, alpha=False)
            image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    else:
        with Image.open(source) as original:
            if original.mode == "RGBA" or "transparency" in original.info:
                rgba = original.convert("RGBA")
                image = Image.new("RGB", rgba.size, "white")
                image.paste(rgba, mask=rgba.getchannel("A"))
            else:
                image = original.convert("RGB")
    with image:
        image.save(out, format="JPEG", quality=95, subsampling=0, dpi=(dpi, dpi))


def _aggregate_breakers(rows: list[dict]) -> list[dict]:
    groups = defaultdict(Decimal)
    for row in rows:
        groups[(row["panel"], row["ampere"], row["poles"])] += number(row["quantity"], "quantity")
    return [{"Panel": panel, "Ampere": amp, "Poles": poles, "Quantity": fmt(q)}
            for (panel, amp, poles), q in sorted(groups.items())]


def _workbook(rows: list[dict], reserves: str, path: Path) -> None:
    book = Workbook()
    book.remove(book.active)

    def tab(title: str, columns: list[str], records: list[list]) -> None:
        ws = book.create_sheet(title)
        ws.append(columns)
        for record in records:
            ws.append(record)
        for row in ws:
            for cell in row:
                # Les données libres ne deviennent jamais des formules Excel.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(70, max(12, max(len(str(v.value or "")) for v in col) + 2))

    purchase = _aggregate_breakers(rows)
    tab("Purchase list", c.BREAKERS, [[r[k] for k in c.BREAKERS] for r in purchase])
    tab("All breakers", c.DETAILED_BREAKERS, [[r[k] for k in c.DETAILED_BREAKERS] for r in rows])
    # Titres déterministes, sans collision avec les feuilles fixes ni troncature silencieuse.
    for index, panel in enumerate(sorted({r["panel"] for r in rows}), 1):
        selected = [r for r in rows if r["panel"] == panel]
        tab(f"P{index:03d}", c.PANEL_BREAKERS, [
            [r["circuits"], r["ampere"], r["poles"], r["quantity"], r["description"], r["protection_source"], r["sheet"]]
            for r in selected])
    tab("Panel index", ["Onglet", "Panneau"], [[f"P{i:03d}", p] for i, p in enumerate(sorted({r["panel"] for r in rows}), 1)])
    tab("Before ordering", ["Informations à confirmer avant commande"], [[line] for line in reserves.splitlines() if line.strip()])
    book.save(path)
    book.close()


def granby(job: Job, out: Path) -> None:
    for sheet in job.sheets:
        if sheet.annotation is None:
            raise ContractError(f"{sheet.id}: annotation requise pour Granby; aucun plan n'est inventé")
        _jpeg(sheet.annotation, out / f"{job.project}-{sheet.id}-releve.jpg", job.dpi)
        for suffix, columns, rows in sheet.tables:
            write_csv(out / f"{job.project}-{sheet.id}-{suffix}.csv", columns, rows)
    if job.breakers:
        write_csv(out / f"{job.project}-breakers-detailed.csv", c.DETAILED_BREAKERS, job.breakers)
        _workbook(job.breakers, job.reserves, out / f"{job.project}-breakers-by-panel.xlsx")
        for sheet in job.sheets:
            selected = [r for r in job.breakers if r["sheet"] == sheet.id]
            if selected:
                write_csv(out / f"{job.project}-{sheet.id}-breakers.csv", c.BREAKERS, _aggregate_breakers(selected))


def chum(job: Job, out: Path) -> None:
    if not job.equipment:
        raise ContractError("CHUM: équipement consolidé requis; metadata obligatoire avec adapter=pipeline")
    for sheet in job.sheets:
        rows = [{
            # Le CSV 4 colonnes ne possède pas Action ni Unité: les garder explicitement dans Type.
            "Type": f"{r['Code']} | {r['Action']} | {r['Unité']}",
            "Description": r["Description"], "Quantité": r["Quantité"], "Source": r["Ref feuille"],
        } for r in job.equipment if r["Plan"] == sheet.id]
        write_csv(out / f"{job.project}-{sheet.id}-equipment.csv", c.EQUIPMENT, rows)
    write_csv(out / f"{job.project}-equipment-consolidated.csv", c.CONSOLIDATED, job.equipment)
    groups = defaultdict(lambda: {"un": Decimal(0), "Lot": Decimal(0)})
    for row in job.equipment:
        groups[(row["Discipline"], row["Action"])][row["Unité"]] += number(row["Quantité"], "Quantité")
    summary = [{"Discipline": d, "Action": a, "Qte (un)": fmt(t["un"]), "Lots": fmt(t["Lot"])}
               for (d, a), t in sorted(groups.items())]
    write_csv(out / f"{job.project}-summary.csv", c.SUMMARY, summary)


def export(job: Job, profile: str, destination: Path) -> list[dict]:
    if profile not in {"granby", "chum"}:
        raise ContractError(f"Profil inconnu: {profile}")
    destination = destination.absolute()
    if destination.exists():
        raise ContractError(f"Destination déjà présente (aucun écrasement): {destination}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".estimation-v2-", dir=destination.parent) as temp:
        out = Path(temp) / "package"
        out.mkdir()
        (granby if profile == "granby" else chum)(job, out)
        (out / "reserves.md").write_text(job.reserves + "\n", encoding="utf-8")
        audit = {
            "schema_version": 1, "profile": profile, "project": job.project,
            "sheets": [s.id for s in job.sheets],
            "scope": job.scope, "verification": "incomplète — emballage uniquement, pas validation du relevé",
            # Les chemins absolus privés ne sont pas propagés dans le paquet.
            "inputs": [{"file": p.name, "bytes": p.stat().st_size, "sha256": digest(p)} for p in sorted(job.inputs)],
        }
        (out / "export-report.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        (out / "READ-ME.txt").write_text(
            f"{job.project} — PROFIL {profile.upper()}\n\nPORTÉE\n{job.scope}\n\n"
            "ÉTAT : vérification incomplète. Export déterministe de données fournies; "
            "aucune nouvelle identification, aucun métré automatique, aucune validation contre l'estimateur.\n"
            "Les plans originaux ne sont pas inclus. Les images proviennent uniquement d'annotations fournies.\n"
            "Les tables spécialisées ne sont jamais additionnées entre elles. Les chambres sont UN extrait; "
            "les symboles de schéma ne sont pas une commande finale.\n"
            "Le sommaire CHUM sépare unités et lots, disciplines et actions. CONSERVER n'est pas une pose.\n"
            "Le CSV par feuille CHUM encode Code | Action | Unité dans Type; le consolidé reste l'autorité métier.\n"
            "Disjoncteurs : les calibres vides restent à confirmer. Purchase list n'est pas une autorisation de commande.\n"
            "Réserves : voir reserves.md. SHA256 et tailles : manifest.json (tous les fichiers sauf lui-même).\n"
            "Les CSV conservent le texte source : les ouvrir comme données, jamais exécuter des formules importées.\n"
            "Ces sorties peuvent contenir des données privées. Ne pas les publier dans le dépôt de code.\n",
            encoding="utf-8",
        )
        manifest = [{"file": p.name, "bytes": p.stat().st_size, "sha256": digest(p)}
                    for p in sorted(out.iterdir())]
        (out / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        # Le dossier n'apparaît qu'après succès de tous les exports et du manifeste.
        os.rename(out, destination)
    return manifest
