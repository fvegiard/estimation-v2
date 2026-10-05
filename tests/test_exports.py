from __future__ import annotations

import csv
import hashlib
import importlib.util
import json
import subprocess
import sys
from decimal import Decimal
from pathlib import Path

import openpyxl
import pymupdf
import pytest
from PIL import Image, ImageChops, ImageStat

from estimation_v2 import contracts as c
from estimation_v2.adapters import load_job
from estimation_v2.cli import main
from estimation_v2.exporters import export
from estimation_v2.io import ContractError, name, read_csv, write_csv

spec = importlib.util.spec_from_file_location("synthetic", Path(__file__).resolve().parents[1] / "examples" / "synthetic.py")
synthetic = importlib.util.module_from_spec(spec)
spec.loader.exec_module(synthetic)


@pytest.fixture
def job_path(tmp_path):
    return synthetic.build(tmp_path / "inputs")


def config(path, transform):
    data = json.loads(path.read_text(encoding="utf-8"))
    transform(data)
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")


def check_manifest(out):
    manifest = json.loads((out / "manifest.json").read_text(encoding="utf-8"))
    assert {r["file"] for r in manifest} == {p.name for p in out.iterdir()} - {"manifest.json"}
    for record in manifest:
        assert list(record) == ["file", "bytes", "sha256"]
        content = (out / record["file"]).read_bytes()
        assert record["bytes"] == len(content)
        assert record["sha256"] == hashlib.sha256(content).hexdigest()


def test_granby_all_contracts_and_source_unchanged(job_path, tmp_path):
    job = load_job(job_path)
    before = {p: p.read_bytes() for p in job.inputs}
    out = tmp_path / "granby"
    export(job, "granby", out)
    assert len(list(out.glob("*.jpg"))) == 8
    expected = {
        "DEMO-E001-equipment.csv": c.EQUIPMENT,
        "DEMO-E002-equipment.csv": ["Type", "Description", "Niveau fictif 1", "Niveau fictif 2", "Total"],
        "DEMO-E003-equipment.csv": c.POSITIONS,
        "DEMO-E004-symbols.csv": c.SYMBOLS,
        "DEMO-E005-by-room-type.csv": c.ROOMS,
        "DEMO-E006-details.csv": c.DETAILS,
        "DEMO-E007-breakers.csv": c.BREAKERS,
        "DEMO-E008-breakers.csv": c.BREAKERS,
        "DEMO-breakers-detailed.csv": c.DETAILED_BREAKERS,
    }
    for filename, columns in expected.items():
        assert read_csv(out / filename, columns)
    assert read_csv(out / "DEMO-E007-breakers.csv", c.BREAKERS)[0]["Quantity"] == "3"
    assert read_csv(out / "DEMO-E008-breakers.csv", c.BREAKERS)[0]["Ampere"] == ""
    book = openpyxl.load_workbook(out / "DEMO-breakers-by-panel.xlsx", read_only=True, data_only=False)
    assert book.sheetnames == ["Purchase list", "All breakers", "P001", "P002", "Panel index", "Before ordering"]
    assert list(book["P001"].values)[0] == tuple(c.PANEL_BREAKERS)
    assert list(book["Panel index"].values)[1:] == [("P001", "P-A"), ("P002", "P-B")]
    assert book["All breakers"].max_row == 4
    book.close()
    check_manifest(out)
    assert all(p.read_bytes() == content for p, content in before.items())
    assert job.reserves in (out / "reserves.md").read_text(encoding="utf-8")
    report = json.loads((out / "export-report.json").read_text(encoding="utf-8"))
    assert report["verification"].startswith("incomplète")
    assert all("\\" not in r["file"] and "/" not in r["file"] for r in report["inputs"])


def test_jpeg_is_the_annotated_page_not_a_blank_placeholder(job_path, tmp_path):
    job = load_job(job_path)
    out = tmp_path / "granby"
    export(job, "granby", out)
    with pymupdf.open(job_path.parent / "annotated.pdf") as pdf:
        pix = pdf[0].get_pixmap(dpi=144, colorspace=pymupdf.csRGB, alpha=False)
        expected = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    with Image.open(out / "DEMO-E001-releve.jpg") as image:
        assert image.size == (840, 600)
        assert image.mode == "RGB"
        assert max(ImageStat.Stat(ImageChops.difference(image, expected)).mean) < 2
        # Zone des trois symboles marqués, distincte du fond blanc.
        for x in [70, 180, 290]:
            region = image.crop((int((x - 5) * 2), 170, int((x + 5) * 2), 190))
            assert min(ImageStat.Stat(region).mean) < 245
    expected.close()


def test_chum_summary_units_lots_actions_and_exact_rows(job_path, tmp_path):
    job = load_job(job_path)
    out = tmp_path / "chum"
    export(job, "chum", out)
    actual = read_csv(out / "DEMO-equipment-consolidated.csv", c.CONSOLIDATED)
    assert actual == job.equipment
    summary = read_csv(out / "DEMO-summary.csv", c.SUMMARY)
    by = {(r["Discipline"], r["Action"]): r for r in summary}
    assert by[("Éclairage", "INSTALLER")]["Qte (un)"] == "3.5"
    assert by[("Éclairage", "CONSERVER")]["Qte (un)"] == "4"
    assert by[("Mise en service", "INSTALLER")]["Lots"] == "2"
    assert by[("Mise en service", "INSTALLER")]["Qte (un)"] == "0"
    assert len(list(out.glob("*-equipment.csv"))) == 8
    assert read_csv(out / "DEMO-E003-equipment.csv", c.EQUIPMENT) == []
    rows = read_csv(out / "DEMO-E001-equipment.csv", c.EQUIPMENT)
    assert rows[0]["Type"] == "DEL-A | INSTALLER | un"
    assert len(rows) == 3
    assert not list(out.glob("*.jpg"))
    check_manifest(out)


def test_no_overwrite_or_partial_success(job_path, tmp_path):
    job = load_job(job_path)
    out = tmp_path / "chum"
    export(job, "chum", out)
    original = {p.name: p.read_bytes() for p in out.iterdir()}
    with pytest.raises(ContractError, match="aucun écrasement"):
        export(job, "chum", out)
    assert original == {p.name: p.read_bytes() for p in out.iterdir()}
    job.sheets[2].annotation["page"] = 99
    failed = tmp_path / "failed"
    with pytest.raises(ContractError, match="absente"):
        export(job, "granby", failed)
    assert not failed.exists()
    assert not list(tmp_path.glob(".estimation-v2-*"))


@pytest.mark.parametrize("value", ["../secret", "a/b", "a\\b", "NUL", "con.txt", "COM1", "a.", "", "é", "x" * 81])
def test_nonportable_names_rejected(value):
    with pytest.raises(ContractError):
        name(value)


@pytest.mark.parametrize("value", ["-1", "NaN", "Infinity", "abc", ""])
def test_invalid_quantities_rejected(job_path, value):
    rows = read_csv(job_path.parent / "equipment.csv", c.CONSOLIDATED)
    rows[0]["Quantité"] = value
    write_csv(job_path.parent / "equipment.csv", c.CONSOLIDATED, rows)
    with pytest.raises(ContractError):
        load_job(job_path)


@pytest.mark.parametrize("col,value", [("Action", "ACHETER"), ("Unité", "m"), ("Plan", "INCONNU"), ("Ref feuille", "")])
def test_invalid_business_fields_rejected(job_path, col, value):
    rows = read_csv(job_path.parent / "equipment.csv", c.CONSOLIDATED)
    rows[0][col] = value
    write_csv(job_path.parent / "equipment.csv", c.CONSOLIDATED, rows)
    with pytest.raises(ContractError):
        load_job(job_path)


def test_level_totals_never_corrected_silently(job_path):
    path = job_path.parent / "table-levels.csv"
    columns = ["Type", "Description", "Niveau fictif 1", "Niveau fictif 2", "Total"]
    rows = read_csv(path, columns)
    rows[0]["Total"] = "900"
    write_csv(path, columns, rows)
    with pytest.raises(ContractError, match="Total"):
        load_job(job_path)


def test_duplicate_breaker_circuit_rejected(job_path):
    path = job_path.parent / "breakers.csv"
    rows = read_csv(path, c.DETAILED_BREAKERS)
    rows.append(dict(rows[0]))
    write_csv(path, c.DETAILED_BREAKERS, rows)
    with pytest.raises(ContractError, match="dupliquée"):
        load_job(job_path)


def test_duplicate_equipment_rejected_but_empty_optional_values_allowed(job_path):
    path = job_path.parent / "equipment.csv"
    rows = read_csv(path, c.CONSOLIDATED)
    assert load_job(job_path).equipment
    rows.append(dict(rows[0]))
    write_csv(path, c.CONSOLIDATED, rows)
    with pytest.raises(ContractError, match="dupliquées"):
        load_job(job_path)


def test_formula_like_text_stays_text_in_workbook(job_path, tmp_path):
    path = job_path.parent / "breakers.csv"
    rows = read_csv(path, c.DETAILED_BREAKERS)
    rows[0]["description"] = '=HYPERLINK("https://invalid.example","ne pas exécuter")'
    write_csv(path, c.DETAILED_BREAKERS, rows)
    out = tmp_path / "granby"
    export(load_job(job_path), "granby", out)
    book = openpyxl.load_workbook(out / "DEMO-breakers-by-panel.xlsx")
    assert book["All breakers"]["F2"].data_type == "s"
    assert book["All breakers"]["F2"].value == rows[0]["description"]
    book.close()


def pipeline_job(job_path):
    directory = job_path.parent / "vecteur"
    directory.mkdir()
    row = dict(zip(c.BORDEREAU, ["E001", "M01-01", "SOURCE-FICTIVE", "MATERIEL-FICTIF",
                               "Designation fictive", "2", "Portee fictive", "Modele fictif",
                               "Identification à revalider", "", "1"]))
    # Colonnes supplémentaires effectivement émises par from_releve.
    row.update({"ref": "Note fictive", "code": "M01", "format": "materiel", "note": "Symbole fictif"})
    write_csv(directory / "bordereau.csv", [*c.BORDEREAU, "ref", "code", "format", "note"], [row])
    estimate = {"source": "releve", "sheets": [{"sheet": "E001", "page": 1, "width_px": 420, "height_px": 300}],
                "counters": [{"family": "FICTIF", "name": "FICTIF", "elements": [
                    {"sheet": "E001", "page": 1, "repere": "M01-01", "code": "M01", "x": 70, "y": 90}]}]}
    (directory / "estimate.json").write_text(json.dumps(estimate), encoding="utf-8")
    (directory / "reserves.md").write_text("## E001\nRéserve originale fictive.", encoding="utf-8")
    meta = dict(zip(c.METADATA, ["E001", "M01-01", "Éclairage", "Local fictif", "un", "INSTALLER",
                               "Pose seule, fourni client", "Note originale fictive"]))
    write_csv(job_path.parent / "metadata.csv", c.METADATA, [meta])
    config(job_path, lambda data: data.update({
        "adapter": {"type": "pipeline", "directory": "vecteur", "metadata": "metadata.csv"},
        "sheets": [{"id": "E001", "annotation": {"path": "annotated.pdf", "page": 1}}],
        "breakers_detail": "breakers.csv",
    }))
    config(job_path, lambda data: data.pop("breakers_detail"))
    return directory


def test_pipeline_adapter_real_output_schema_without_source_dependency(job_path, tmp_path):
    pipeline_job(job_path)
    job = load_job(job_path)
    assert job.sheets[0].tables[0][2][0]["X points"] == "70"
    assert job.equipment[0]["Quantité"] == "2"
    assert job.equipment[0]["Règle appliquée"] == "Pose seule, fourni client"
    assert "Réserve originale fictive." in job.reserves
    assert "E001 / M01-01" in job.reserves
    export(job, "granby", tmp_path / "granby")
    export(job, "chum", tmp_path / "chum")
    assert read_csv(tmp_path / "granby" / "DEMO-E001-equipment.csv", c.POSITIONS)[0]["Repère"] == "M01-01"
    assert read_csv(tmp_path / "chum" / "DEMO-summary.csv", c.SUMMARY)[0]["Qte (un)"] == "2"


@pytest.mark.parametrize("mutation", ["duplicate", "missing", "orphan", "outside", "pixels", "wrongpage"])
def test_pipeline_join_and_coordinates_strict(job_path, mutation):
    directory = pipeline_job(job_path)
    path = directory / "estimate.json"
    est = json.loads(path.read_text(encoding="utf-8"))
    elements = est["counters"][0]["elements"]
    if mutation == "duplicate":
        elements.append(dict(elements[0]))
    elif mutation == "missing":
        elements.clear()
    elif mutation == "orphan":
        elements[0]["repere"] = "UNKNOWN"
    elif mutation == "outside":
        elements[0]["x"] = 421
    elif mutation == "pixels":
        est["source"] = "detector"
    else:
        elements[0]["page"] = 2
    path.write_text(json.dumps(est), encoding="utf-8")
    with pytest.raises(ContractError):
        load_job(job_path)


def test_pipeline_metadata_required_only_for_chum(job_path, tmp_path):
    pipeline_job(job_path)
    config(job_path, lambda data: data["adapter"].pop("metadata"))
    job = load_job(job_path)
    export(job, "granby", tmp_path / "granby")
    with pytest.raises(ContractError, match="metadata obligatoire"):
        export(job, "chum", tmp_path / "chum")
    assert not (tmp_path / "chum").exists()


def test_pipeline_metadata_exhaustive(job_path):
    pipeline_job(job_path)
    write_csv(job_path.parent / "metadata.csv", c.METADATA, [])
    with pytest.raises(ContractError, match="exhaustive"):
        load_job(job_path)


def test_missing_annotation_does_not_invent_a_plan(job_path, tmp_path):
    config(job_path, lambda data: data["sheets"][0].pop("annotation"))
    with pytest.raises(ContractError, match="aucun plan"):
        export(load_job(job_path), "granby", tmp_path / "granby")


def test_unknown_schema_field_is_not_ignored(job_path):
    config(job_path, lambda data: data.update({"guess_quantities": True}))
    with pytest.raises(ContractError, match="inconnues"):
        load_job(job_path)


def test_cli_success_failure_and_standalone_module(job_path, tmp_path, capsys):
    assert main(["--job", str(job_path), "--profile", "chum", "--out", str(tmp_path / "chum")]) == 0
    assert main(["--job", str(job_path), "--profile", "chum", "--out", str(tmp_path / "chum")]) == 2
    assert "ÉCHEC" in capsys.readouterr().err
    result = subprocess.run([sys.executable, "-m", "estimation_v2", "--help"], capture_output=True)
    assert result.returncode == 0


def test_csv_width_and_header_errors(tmp_path):
    path = tmp_path / "invalid.csv"
    path.write_text("Type,Description,Quantité,Source\nA,B,1,S,EXTRA\n", encoding="utf-8")
    with pytest.raises(ContractError, match="largeur"):
        read_csv(path, c.EQUIPMENT)
    path.write_text("bad,columns\n", encoding="utf-8")
    with pytest.raises(ContractError, match="colonnes"):
        read_csv(path, c.EQUIPMENT)


def test_both_packages_repeatable_at_logical_level(job_path, tmp_path):
    job = load_job(job_path)
    for profile in ["granby", "chum"]:
        a, b = tmp_path / (profile + "-a"), tmp_path / (profile + "-b")
        export(job, profile, a)
        export(job, profile, b)
        for path in a.iterdir():
            if path.suffix in {".csv", ".txt", ".md", ".jpg"}:
                assert path.read_bytes() == (b / path.name).read_bytes()
